#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0.2", "PyYAML>=6.0.2"]
# ///

"""Convert one LENS Markdown paper note into a figure-led PPTX."""

from __future__ import annotations

import argparse
import configparser
import html
import os
import re
import tempfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterable

import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


DEFAULT_EXCLUDED = ("Works", "Question–Method Map", "Key references")
DEFAULT_CONFIG = Path(__file__).resolve().parents[1] / "ppt.config"
HEADING_RE = re.compile(r"^(#{1,4})\s+(.+?)\s*$")
CENTER_RE = re.compile(r"<center\b[^>]*>(.*?)</center>", re.I | re.S)
IMG_RE = re.compile(r"<img\b[^>]*\bsrc\s*=\s*(['\"])(.*?)\1", re.I | re.S)
DIV_RE = re.compile(r"<div\b[^>]*>(.*?)</div>", re.I | re.S)
MARKDOWN_IMAGE_RE = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")
VISUAL_PREFIX_RE = re.compile(
    r"^\s*(?:Fig(?:ure)?\.?|Table)\s*\d+[A-Za-z]?\s*[.：:\-]?\s*",
    re.I,
)


@dataclass
class ImageSpec:
    source: str
    caption: str = ""
    references: list[str] = field(default_factory=list)
    path: Path | None = None


@dataclass
class Block:
    level: int
    title: str
    lines: list[str] = field(default_factory=list)
    children: list["Block"] = field(default_factory=list)


@dataclass
class SlideSpec:
    kind: str
    title: str
    subtitle: str = ""
    images: list[ImageSpec] = field(default_factory=list)
    references: list[str] = field(default_factory=list)
    items: list[str] = field(default_factory=list)
    active_index: int | None = None
    image_legends_above: bool = False


def normalize_heading(value: str) -> str:
    value = value.replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", value).strip().casefold()


def parse_frontmatter(text: str) -> tuple[dict, str]:
    match = re.match(r"^---\s*\n(.*?)\n---\s*(?:\n|$)", text, re.S)
    if not match:
        return {}, text
    try:
        metadata = yaml.safe_load(match.group(1)) or {}
        if not isinstance(metadata, dict):
            metadata = {}
    except yaml.YAMLError:
        metadata = {}
    return metadata, text[match.end() :]


def parse_blocks(markdown: str, excluded: set[str]) -> list[Block]:
    roots: list[Block] = []
    stack: list[Block] = []
    excluded_h1 = False
    for raw_line in markdown.replace("\r\n", "\n").split("\n"):
        match = HEADING_RE.match(raw_line.strip())
        if match:
            level = len(match.group(1))
            title = match.group(2).strip()
            if level == 1:
                excluded_h1 = normalize_heading(title) in excluded
                stack = []
                if excluded_h1:
                    continue
            elif excluded_h1:
                continue
            while stack and stack[-1].level >= level:
                stack.pop()
            block = Block(level=level, title=title)
            if stack:
                stack[-1].children.append(block)
            else:
                roots.append(block)
            stack.append(block)
        elif not excluded_h1 and stack:
            stack[-1].lines.append(raw_line)
    return roots


def strip_html(value: str) -> str:
    value = re.sub(r"<br\s*/?>", " ", value, flags=re.I)
    value = re.sub(r"<[^>]+>", "", value)
    return re.sub(r"\s+", " ", html.unescape(value)).strip()


def adjacent_references(source: list[str], image_end_line: int) -> list[str]:
    """Return only blockquotes directly adjacent to one image block."""
    next_index = image_end_line + 1
    while next_index < len(source) and re.match(
        r"^\s*<!--\s*(?:figure|table):", source[next_index], re.I
    ):
        next_index += 1
    if next_index >= len(source) or not source[next_index].strip():
        return []
    references: list[str] = []
    while next_index < len(source):
        stripped = source[next_index].strip()
        if not stripped.startswith(">") or not stripped[1:].strip():
            break
        references.append(stripped[1:].strip())
        next_index += 1
    return references


def extract_image_entries(
    lines: Iterable[str],
) -> tuple[list[str], list[tuple[int, int, ImageSpec]]]:
    source = list(lines)
    text = "\n".join(source)
    entries: list[tuple[int, int, ImageSpec]] = []
    occupied: list[tuple[int, int]] = []
    for center in CENTER_RE.finditer(text):
        body = center.group(1)
        image_match = IMG_RE.search(body)
        if not image_match:
            continue
        caption_match = DIV_RE.search(body)
        start_line = text.count("\n", 0, center.start())
        end_line = text.count("\n", 0, center.end())
        entries.append(
            (
                start_line,
                end_line,
                ImageSpec(
                    source=html.unescape(image_match.group(2).strip()),
                    caption=(
                        strip_html(caption_match.group(1)) if caption_match else ""
                    ),
                    references=adjacent_references(source, end_line),
                ),
            )
        )
        occupied.append(center.span())
    for match in MARKDOWN_IMAGE_RE.finditer(text):
        if not any(start <= match.start() < end for start, end in occupied):
            start_line = text.count("\n", 0, match.start())
            end_line = text.count("\n", 0, match.end())
            entries.append(
                (
                    start_line,
                    end_line,
                    ImageSpec(
                        source=html.unescape(match.group(1).strip()),
                        references=adjacent_references(source, end_line),
                    ),
                )
            )
    entries.sort(key=lambda entry: (entry[0], entry[1]))
    return source, entries


def extract_images(lines: Iterable[str]) -> list[ImageSpec]:
    _, entries = extract_image_entries(lines)
    return [image for _, _, image in entries]


def image_unit_end_line(source: list[str], image_end_line: int) -> int:
    """Return the last line belonging to an image and its adjacent Citation."""
    cursor = image_end_line + 1
    while cursor < len(source) and re.match(
        r"^\s*<!--\s*(?:figure|table):", source[cursor], re.I
    ):
        cursor += 1
    last_line = cursor - 1
    while cursor < len(source):
        stripped = source[cursor].strip()
        if not stripped.startswith(">") or not stripped[1:].strip():
            break
        last_line = cursor
        cursor += 1
    return max(image_end_line, last_line)


def extract_contiguous_images(lines: Iterable[str]) -> list[ImageSpec]:
    """Keep only the first uninterrupted image/Citation sequence."""
    source, entries = extract_image_entries(lines)
    if not entries:
        return []
    _, end_line, first_image = entries[0]
    selected = [first_image]
    unit_end = image_unit_end_line(source, end_line)
    for start_line, end_line, image in entries[1:]:
        if start_line != unit_end + 1:
            break
        selected.append(image)
        unit_end = image_unit_end_line(source, end_line)
    return selected


def extract_references(lines: Iterable[str]) -> list[str]:
    return [
        reference
        for image in extract_images(lines)
        for reference in image.references
    ]


def subtree_lines(block: Block) -> list[str]:
    result = list(block.lines)
    for child in block.children:
        result.extend(subtree_lines(child))
    return result


def cleaned_visual_title(caption: str, fallback: str) -> str:
    title = VISUAL_PREFIX_RE.sub("", strip_html(caption)).strip()
    title = re.sub(r"[。.]\s*$", "", title).strip()
    if title:
        return title
    title = VISUAL_PREFIX_RE.sub("", fallback).strip()
    return re.sub(r"[。.]\s*$", "", title).strip() or fallback


def first_h4_title(block: Block) -> str:
    return next((child.title for child in block.children if child.level == 4), "")


def content_spec(block: Block) -> SlideSpec:
    lines = subtree_lines(block)
    images = (
        extract_contiguous_images(lines)
        if block.level == 3
        else extract_images(lines)
    )
    return SlideSpec(
        kind="content",
        title=block.title,
        subtitle=first_h4_title(block),
        images=images,
        references=[
            reference for image in images for reference in image.references
        ],
        image_legends_above=block.level == 3,
    )


def visual_block_title(block: Block) -> str:
    direct_images = extract_images(block.lines)
    caption = direct_images[0].caption if direct_images else ""
    return cleaned_visual_title(caption, block.title)


def append_h2_content(slides: list[SlideSpec], h2: Block, *, divider: bool) -> None:
    h3_children = [child for child in h2.children if child.level == 3]
    direct_images = extract_images(h2.lines)
    if direct_images:
        slides.append(
            SlideSpec(
                kind="content",
                title=cleaned_visual_title(direct_images[0].caption, h2.title),
                subtitle=first_h4_title(h2),
                images=direct_images,
                references=extract_references(h2.lines),
            )
        )
        slides.extend(content_spec(h3) for h3 in h3_children)
        return
    if h3_children:
        if divider:
            slides.append(SlideSpec(kind="divider", title=h2.title))
        slides.extend(content_spec(h3) for h3 in h3_children)
        return
    lines = subtree_lines(h2)
    images = extract_images(lines)
    if images:
        slides.append(
            SlideSpec(
                kind="content",
                title=cleaned_visual_title(images[0].caption, h2.title),
                subtitle=first_h4_title(h2),
                images=images,
                references=extract_references(lines),
            )
        )


def order_top_level_sections(roots: list[Block]) -> list[Block]:
    preferred = ("info", "intro", "summary")
    ordered: list[Block] = []
    consumed: set[int] = set()
    for target in preferred:
        for index, block in enumerate(roots):
            if index not in consumed and normalize_heading(block.title) == target:
                ordered.append(block)
                consumed.add(index)
                break
    ordered.extend(block for index, block in enumerate(roots) if index not in consumed)
    return ordered


def collect_slide_specs(roots: list[Block]) -> list[SlideSpec]:
    slides: list[SlideSpec] = []
    ordered_roots = order_top_level_sections(roots)
    results_block = next(
        (h1 for h1 in ordered_roots if normalize_heading(h1.title) == "results"),
        None,
    )
    results_h2 = (
        [child for child in results_block.children if child.level == 2]
        if results_block
        else []
    )
    results_items = [visual_block_title(h2) for h2 in results_h2]
    overview_inserted = False
    for h1 in ordered_roots:
        if h1.level != 1:
            continue
        slides.append(SlideSpec(kind="divider", title=h1.title))
        h2_children = [child for child in h1.children if child.level == 2]
        if normalize_heading(h1.title) == "results" and h2_children:
            if not overview_inserted:
                slides.append(
                    SlideSpec(
                        kind="results_outline",
                        title="Overall study strategy and workflow",
                        items=results_items,
                    )
                )
                overview_inserted = True
            for index, h2 in enumerate(h2_children):
                slides.append(
                    SlideSpec(
                        kind="results_outline",
                        title="Overall study strategy and workflow",
                        items=results_items,
                        active_index=index,
                    )
                )
                append_h2_content(slides, h2, divider=False)
        else:
            for h2 in h2_children:
                append_h2_content(slides, h2, divider=True)
        slides.extend(
            content_spec(h3) for h3 in h1.children if h3.level == 3
        )
        if (
            normalize_heading(h1.title) == "summary"
            and results_items
            and not overview_inserted
        ):
            slides.append(
                SlideSpec(
                    kind="results_outline",
                    title="Overall study strategy and workflow",
                    items=results_items,
                )
            )
            overview_inserted = True
    return slides


def resolve_image(note_path: Path, asset_dir: Path, image: ImageSpec) -> Path | None:
    source = image.source.strip().strip("<>")
    if source.startswith(("http://", "https://")):
        return None
    raw = Path(source).expanduser()
    candidates = [raw] if raw.is_absolute() else [
        note_path.parent / raw,
        asset_dir / raw,
        asset_dir / raw.name,
    ]
    return next((path.resolve() for path in candidates if path.is_file()), None)


def load_ppt_config() -> configparser.ConfigParser:
    config_path = Path(
        os.environ.get("LENS_NOTE_TO_PPT_CONFIG", str(DEFAULT_CONFIG))
    ).expanduser().resolve()
    if not config_path.is_file():
        raise FileNotFoundError(f"PPT config not found: {config_path}")
    config = configparser.ConfigParser(interpolation=None)
    config.read(config_path, encoding="utf-8")
    return config


def config_value(config, section: str, key: str, fallback):
    if isinstance(fallback, bool):
        return config.getboolean(section, key, fallback=fallback)
    if isinstance(fallback, float):
        return config.getfloat(section, key, fallback=fallback)
    return config.get(section, key, fallback=fallback)


def config_color(config, section: str, fallback: str = "000000") -> RGBColor:
    value = str(config_value(config, section, "font_color", fallback)).strip().lstrip("#")
    if not re.fullmatch(r"[0-9A-Fa-f]{6}", value):
        raise ValueError(f"Invalid font_color in [{section}]: {value}")
    return RGBColor.from_string(value.upper())


def config_align(config, section: str, fallback: str = "center"):
    value = str(config_value(config, section, "align", fallback)).strip().casefold()
    return {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
    }.get(value, PP_ALIGN.CENTER)


def set_run_font(run, config, section: str) -> None:
    run.font.size = Pt(config_value(config, section, "font_size_pt", 18.0))
    run.font.bold = config_value(config, section, "bold", False)
    run.font.italic = config_value(config, section, "italic", False)
    run.font.name = config_value(config, section, "font_name", "Arial")
    properties = run._r.get_or_add_rPr()
    east_asian = properties.find(qn("a:ea"))
    if east_asian is None:
        east_asian = OxmlElement("a:ea")
        properties.append(east_asian)
    east_asian.set(
        "typeface", config_value(config, section, "east_asia_font", "SimHei")
    )
    run.font.color.rgb = config_color(config, section)


def set_text(box, text: str, config, section: str) -> None:
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = config_align(config, section)
    run = paragraph.add_run()
    run.text = text
    set_run_font(run, config, section)


def add_page_numbers(prs: Presentation, config) -> None:
    if not config_value(config, "page_number", "show", True):
        return
    show_on_cover = config_value(config, "page_number", "show_on_cover", False)
    first_numbered_index = 0 if show_on_cover else 1
    start_at = int(config_value(config, "page_number", "start_at", "1"))
    width = Inches(config_value(config, "page_number", "width_in", 0.36))
    height = Inches(config_value(config, "page_number", "height_in", 0.24))
    right = Inches(config_value(config, "page_number", "right_in", 0.18))
    bottom = Inches(config_value(config, "page_number", "bottom_in", 0.10))
    for slide_index, slide in enumerate(prs.slides):
        if slide_index < first_numbered_index:
            continue
        page_number = start_at + slide_index - first_numbered_index
        box = slide.shapes.add_textbox(
            prs.slide_width - right - width,
            prs.slide_height - bottom - height,
            width,
            height,
        )
        set_text(box, str(page_number), config, "page_number")


def add_cover(prs: Presentation, title: str, reference: str, presenter: str, config) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    sw, sh = prs.slide_width, prs.slide_height
    kicker_left = Inches(config_value(config, "cover_kicker", "left_in", 1.0))
    kicker_right = Inches(config_value(config, "cover_kicker", "right_in", 1.0))
    kicker = slide.shapes.add_textbox(
        kicker_left,
        Inches(config_value(config, "cover_kicker", "top_in", 2.15)),
        sw - kicker_left - kicker_right,
        Inches(config_value(config, "cover_kicker", "height_in", 0.65)),
    )
    set_text(kicker, config_value(config, "cover_kicker", "text", "Journal Report"), config, "cover_kicker")
    title_left = Inches(config_value(config, "paper_title", "left_in", 1.0))
    title_right = Inches(config_value(config, "paper_title", "right_in", 1.0))
    box = slide.shapes.add_textbox(
        title_left,
        Inches(config_value(config, "paper_title", "top_in", 2.95)),
        sw - title_left - title_right,
        Inches(config_value(config, "paper_title", "height_in", 1.70)),
    )
    set_text(box, title, config, "paper_title")
    info_width = Inches(config_value(config, "presenter", "width_in", 4.0))
    info = slide.shapes.add_textbox(
        sw - Inches(config_value(config, "presenter", "right_in", 0.75)) - info_width,
        sh - Inches(config_value(config, "presenter", "bottom_in", 0.70)) - Inches(config_value(config, "presenter", "height_in", 0.45)),
        info_width,
        Inches(config_value(config, "presenter", "height_in", 0.45)),
    )
    set_text(info, f"{presenter} {datetime.now():%Y-%m-%d}", config, "presenter")
    if reference:
        footer = citation_box(slide, sw, sh, config)
        set_text(footer, reference, config, "citation")


def add_divider(prs: Presentation, title: str, config) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    sw, sh = prs.slide_width, prs.slide_height
    left = Inches(config_value(config, "section_title", "left_in", 1.0))
    right = Inches(config_value(config, "section_title", "right_in", 1.0))
    height = Inches(config_value(config, "section_title", "height_in", 0.78))
    box = slide.shapes.add_textbox(left, (sh - height) // 2, sw - left - right, height)
    set_text(box, title, config, "section_title")


def add_results_outline(prs: Presentation, spec: SlideSpec, config) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    sw, sh = prs.slide_width, prs.slide_height
    title_left = Inches(config_value(config, "results_outline_title", "left_in", 0.70))
    title_right = Inches(config_value(config, "results_outline_title", "right_in", 0.70))
    title = slide.shapes.add_textbox(
        title_left,
        Inches(config_value(config, "results_outline_title", "top_in", 0.22)),
        sw - title_left - title_right,
        Inches(config_value(config, "results_outline_title", "height_in", 0.65)),
    )
    set_text(
        title,
        config_value(
            config,
            "results_outline_title",
            "text",
            spec.title or "Overall study strategy and workflow",
        ),
        config,
        "results_outline_title",
    )
    if not spec.items:
        return
    list_left = Inches(config_value(config, "results_outline_item", "left_in", 1.00))
    list_right = Inches(config_value(config, "results_outline_item", "right_in", 0.50))
    list_top = Inches(config_value(config, "results_outline_item", "list_top_in", 1.25))
    list_bottom = Inches(config_value(config, "results_outline_item", "list_bottom_in", 0.60))
    gap = Inches(config_value(config, "results_outline_item", "gap_in", 0.12))
    available = sh - list_top - list_bottom
    item_height = min(
        Inches(config_value(config, "results_outline_item", "height_in", 0.50)),
        (available - gap * (len(spec.items) - 1)) // len(spec.items),
    )
    group_height = item_height * len(spec.items) + gap * (len(spec.items) - 1)
    top = list_top + (available - group_height) // 2
    for index, item in enumerate(spec.items):
        box = slide.shapes.add_textbox(
            list_left,
            top + index * (item_height + gap),
            sw - list_left - list_right,
            item_height,
        )
        style = (
            "results_outline_inactive_item"
            if spec.active_index is not None and index != spec.active_index
            else "results_outline_item"
        )
        set_text(box, f"{index + 1}.    {item}", config, style)


def citation_box(slide, sw: int, sh: int, config):
    left = Inches(config_value(config, "citation", "left_in", 0.50))
    right = Inches(config_value(config, "citation", "right_in", 0.50))
    height = Inches(config_value(config, "citation", "height_in", 0.58))
    bottom = Inches(config_value(config, "citation", "bottom_in", 0.04))
    return slide.shapes.add_textbox(left, sh - bottom - height, sw - left - right, height)


def fit_picture(slide, path: Path, left: int, top: int, width: int, height: int):
    picture = slide.shapes.add_picture(str(path), 0, 0)
    scale = min(width / picture.width, height / picture.height)
    picture.width = int(picture.width * scale)
    picture.height = int(picture.height * scale)
    picture.left = left + (width - picture.width) // 2
    picture.top = top + (height - picture.height) // 2
    return picture


def add_content(prs: Presentation, spec: SlideSpec, config) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    sw, sh = prs.slide_width, prs.slide_height
    title_left = Inches(config_value(config, "content_title", "left_in", 0.55))
    title_right = Inches(config_value(config, "content_title", "right_in", 0.55))
    title = slide.shapes.add_textbox(
        title_left,
        Inches(config_value(config, "content_title", "top_in", 0.08)),
        sw - title_left - title_right,
        Inches(config_value(config, "content_title", "height_in", 0.58)),
    )
    set_text(title, spec.title, config, "content_title")
    if spec.subtitle:
        subtitle_left = Inches(config_value(config, "content_subtitle", "left_in", 0.80))
        subtitle_right = Inches(config_value(config, "content_subtitle", "right_in", 0.80))
        subtitle = slide.shapes.add_textbox(
            subtitle_left,
            Inches(config_value(config, "content_subtitle", "top_in", 0.68)),
            sw - subtitle_left - subtitle_right,
            Inches(config_value(config, "content_subtitle", "height_in", 0.42)),
        )
        set_text(subtitle, spec.subtitle, config, "content_subtitle")
    media_top = Inches(1.16) if spec.subtitle else Inches(0.78)
    media_bottom = sh - Inches(0.18)
    selected = [image for image in spec.images if image.path is not None][:2]
    if selected:
        outer_margin = Inches(0.6)
        gap = Inches(0.35)
        count = len(selected)
        cell_width = sw - 2 * outer_margin if count == 1 else (sw - 2 * outer_margin - gap) // 2
        show_image_captions = (
            spec.image_legends_above
            and config_value(config, "image_caption", "show_above_h3", True)
        )
        caption_height = Inches(
            config_value(config, "image_caption", "height_in", 0.45)
        )
        caption_gap = Inches(
            config_value(config, "image_caption", "gap_below_in", 0.08)
        )
        citation_height = Inches(
            config_value(config, "citation", "image_height_in", 0.35)
        )
        citation_gap = Inches(
            config_value(config, "citation", "image_gap_above_in", 0.04)
        )
        for index, image in enumerate(selected):
            left = outer_margin if count == 1 else outer_margin + index * (cell_width + gap)
            has_caption = show_image_captions and bool(image.caption)
            picture_top = (
                media_top + caption_height + caption_gap if has_caption else media_top
            )
            has_citation = bool(image.references)
            picture_bottom = media_bottom - (
                citation_gap + citation_height if has_citation else 0
            )
            picture_height = max(Inches(1.0), picture_bottom - picture_top)
            picture = fit_picture(
                slide, image.path, left, picture_top, cell_width, picture_height
            )
            if has_caption:
                caption = slide.shapes.add_textbox(
                    left, media_top, cell_width, caption_height
                )
                prefix = config_value(config, "image_caption", "prefix", "□")
                set_text(
                    caption,
                    f"{prefix}  {image.caption}" if prefix else image.caption,
                    config,
                    "image_caption",
                )
            if has_citation:
                citation = slide.shapes.add_textbox(
                    left,
                    picture.top + picture.height + citation_gap,
                    cell_width,
                    citation_height,
                )
                set_text(citation, "\n".join(image.references), config, "citation")


def cover_reference(metadata: dict) -> str:
    citation = metadata.get("citation")
    if isinstance(citation, str):
        lines = [line.strip() for line in citation.splitlines() if line.strip()]
        if lines:
            return lines[0]
    return ", ".join(
        value for value in (
            str(metadata.get("journal", "")).strip(),
            str(metadata.get("year", "")).strip(),
        ) if value
    )


def build_presentation(note_path: Path, output_path: Path) -> int:
    config = load_ppt_config()
    metadata, body = parse_frontmatter(note_path.read_text(encoding="utf-8"))
    excluded_raw = os.environ.get(
        "LENS_NOTE_TO_PPT_EXCLUDED_SECTIONS", "|".join(DEFAULT_EXCLUDED)
    )
    excluded = {
        normalize_heading(value) for value in excluded_raw.split("|") if value.strip()
    }
    specs = collect_slide_specs(parse_blocks(body, excluded))
    asset_dir = note_path.parent / "assets" / note_path.stem
    missing_images: list[str] = []
    for spec in specs:
        for image in spec.images:
            image.path = resolve_image(note_path, asset_dir, image)
            if image.path is None:
                missing_images.append(image.source)
    if missing_images:
        raise FileNotFoundError("Missing image assets: " + ", ".join(sorted(set(missing_images))))

    prs = Presentation()
    aspect_ratio = os.environ.get(
        "LENS_NOTE_TO_PPT_ASPECT_RATIO",
        config_value(config, "deck", "aspect_ratio", "16:9"),
    )
    if aspect_ratio == "16:9":
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    else:
        prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    title = str(metadata.get("title") or note_path.stem)
    add_cover(
        prs,
        title,
        cover_reference(metadata),
        os.environ.get("LENS_NOTE_TO_PPT_PRESENTER", "JGH"),
        config,
    )
    for spec in specs:
        if spec.kind == "divider":
            add_divider(prs, spec.title, config)
        elif spec.kind == "results_outline":
            add_results_outline(prs, spec, config)
        else:
            add_content(prs, spec, config)
    add_page_numbers(prs, config)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix=f".{output_path.stem}.", suffix=".pptx", dir=output_path.parent, delete=False
    ) as handle:
        temporary_path = Path(handle.name)
    try:
        prs.save(temporary_path)
        verified = Presentation(temporary_path)
        if not verified.slides or len(verified.slides) != len(prs.slides):
            raise RuntimeError("PPTX verification failed: slide count mismatch")
        os.replace(temporary_path, output_path)
    finally:
        temporary_path.unlink(missing_ok=True)
    return len(prs.slides)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert one LENS Markdown paper note into a figure-led PPTX."
    )
    parser.add_argument("note", type=Path, help="Path to one LENS Markdown note")
    parser.add_argument(
        "--output", type=Path,
        help="Optional output path; defaults to assets/<note-stem>/<note-stem>.pptx",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    note_path = args.note.expanduser().resolve()
    if not note_path.is_file():
        raise SystemExit(f"Note not found: {note_path}")
    if note_path.suffix.lower() != ".md":
        raise SystemExit(f"Expected a Markdown note: {note_path}")
    default_output = note_path.parent / "assets" / note_path.stem / f"{note_path.stem}.pptx"
    output_path = (args.output or default_output).expanduser().resolve()
    if output_path.suffix.lower() != ".pptx":
        raise SystemExit(f"Output must use the .pptx extension: {output_path}")
    slide_count = build_presentation(note_path, output_path)
    print(f"PPT created: {output_path}")
    print(f"Slides: {slide_count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
